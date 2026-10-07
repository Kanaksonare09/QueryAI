"""
Document text extractor — supports PDF, DOCX, TXT, CSV, MD, JSON.
Preserves page numbers and section metadata where possible.
"""
import json
import csv
import io
from pathlib import Path
from typing import Any, Dict, List, Optional
import structlog

logger = structlog.get_logger()


class DocumentExtractor:
    """Extracts text and metadata from various document formats."""

    def extract(self, file_path: str, file_type: str) -> Dict[str, Any]:
        """
        Extract text from a document file.
        Returns: {text, pages, metadata, sections}
        """
        path = Path(file_path)
        file_type = file_type.lower().lstrip(".")

        extractors = {
            "pdf": self._extract_pdf,
            "docx": self._extract_docx,
            "txt": self._extract_txt,
            "md": self._extract_txt,
            "csv": self._extract_csv,
            "json": self._extract_json,
        }

        extractor = extractors.get(file_type)
        if not extractor:
            raise ValueError(f"Unsupported file type: {file_type}")

        return extractor(path)

    def _extract_pdf(self, path: Path) -> Dict[str, Any]:
        import fitz  # PyMuPDF
        doc = fitz.open(str(path))
        pages = []
        full_text_parts = []

        for page_num, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text.strip():
                pages.append({"page": page_num, "text": text})
                full_text_parts.append(f"[Page {page_num}]\n{text}")

        doc.close()
        return {
            "text": "\n\n".join(full_text_parts),
            "pages": pages,
            "page_count": len(pages),
            "metadata": {"source": "pdf"},
            "sections": [],
        }

    def _extract_docx(self, path: Path) -> Dict[str, Any]:
        from docx import Document as DocxDocument
        doc = DocxDocument(str(path))
        sections = []
        full_text_parts = []
        current_section = None

        for para in doc.paragraphs:
            if para.style.name.startswith("Heading"):
                current_section = para.text.strip()
            if para.text.strip():
                if current_section:
                    full_text_parts.append(f"[Section: {current_section}]\n{para.text}")
                    sections.append({"section": current_section, "text": para.text})
                else:
                    full_text_parts.append(para.text)

        # Include tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    full_text_parts.append(row_text)

        return {
            "text": "\n\n".join(full_text_parts),
            "pages": [],
            "page_count": None,
            "metadata": {"source": "docx"},
            "sections": sections,
        }

    def _extract_txt(self, path: Path) -> Dict[str, Any]:
        text = path.read_text(encoding="utf-8", errors="replace")
        return {
            "text": text,
            "pages": [],
            "page_count": None,
            "metadata": {"source": "text"},
            "sections": [],
        }

    def _extract_csv(self, path: Path) -> Dict[str, Any]:
        text_parts = []
        with open(path, newline="", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            rows = list(reader)

        if not rows:
            return {"text": "", "pages": [], "page_count": None, "metadata": {}, "sections": []}

        headers = rows[0]
        text_parts.append("CSV Data:\n" + " | ".join(headers))
        text_parts.append("-" * 40)
        for row in rows[1:]:
            row_dict = dict(zip(headers, row))
            text_parts.append(", ".join(f"{k}: {v}" for k, v in row_dict.items()))

        return {
            "text": "\n".join(text_parts),
            "pages": [],
            "page_count": None,
            "metadata": {"source": "csv", "columns": headers, "row_count": len(rows) - 1},
            "sections": [],
        }

    def _extract_json(self, path: Path) -> Dict[str, Any]:
        data = json.loads(path.read_text(encoding="utf-8"))
        text = json.dumps(data, indent=2)
        return {
            "text": text,
            "pages": [],
            "page_count": None,
            "metadata": {"source": "json"},
            "sections": [],
        }
